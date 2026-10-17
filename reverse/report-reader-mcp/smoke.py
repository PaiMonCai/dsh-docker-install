"""Real converter and MCP stdio acceptance; no model, remote service or sample execution."""
import asyncio
import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile
import zipfile
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from reader import MAX_TEXT, read_document


def pdf_bytes():
    stream = b"BT /F1 12 Tf 72 720 Td (Reader evidence) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:])
    return data + f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


def create_docx(path, content):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", """<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>""")
        archive.writestr("_rels/.rels", """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>""")
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>' + content + '</w:t></w:r></w:p></w:body></w:document>')


def create_documents(root):
    (root / "report.pdf").write_bytes(pdf_bytes())
    create_docx(root / "report.docx", "Reader evidence")
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(["Finding", "Evidence"])
    book.active.append(["Reader evidence", "Recorded"])
    book.save(root / "report.xlsx")
    from pptx import Presentation
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[0])
    slide.shapes.title.text = "Reader evidence"
    slides.save(root / "report.pptx")


async def smoke(root):
    create_documents(root)
    for suffix in ("pdf", "docx", "xlsx", "pptx"):
        path = root / ("report." + suffix)
        before = path.read_bytes()
        print(f"Report reader smoke: converting {suffix}", flush=True)
        try:
            result = read_document(root, str(path), max_chars=12000)
        except Exception as error:
            raise AssertionError(f"report reader {suffix} conversion failed") from error
        print(f"Report reader smoke: {suffix} passed", flush=True)
        assert "Reader evidence" in result["text"], (suffix, result)
        assert result["source"]["sha256"] == hashlib.sha256(before).hexdigest()
        assert path.read_bytes() == before
    first = read_document(root, "report.docx", max_chars=6)
    following = read_document(root, "report.docx", start=first["next_start"])
    assert "Reader evidence" in first["text"] + following["text"]
    create_docx(root / "long.docx", "Evidence " * 33000)
    bounded = read_document(root, "long.docx", max_chars=100)
    assert bounded["conversion_truncated"] and bounded["converted_chars"] == MAX_TEXT
    assert len(bounded["text"]) == 100 and bounded["next_start"] == 100
    (root / "findings.sarif").write_text(json.dumps({
        "version": "2.1.0", "runs": [{"tool": {"driver": {"name": "fixture", "rules": [
            {"id": "one", "defaultConfiguration": {"level": "error"}}
        ]}}, "results": [{"ruleId": "one", "message": {"text": "Reader evidence"}}]}]
    }), encoding="utf-8")
    params = StdioServerParameters(
        command=sys.executable, args=[str(Path(__file__).with_name("server.py"))],
        env={**os.environ, "DSH_REVERSE_WORKSPACE_ROOT": str(root)},
    )
    async with stdio_client(params) as (receive, send):
        async with ClientSession(receive, send) as session:
            await session.initialize()
            tools = (await session.list_tools()).tools
            assert {tool.name for tool in tools} == {"read_sarif", "read_document"}
            assert all(tool.annotations.readOnlyHint and not tool.annotations.openWorldHint for tool in tools)
            result = await session.call_tool("read_sarif", {"path": "findings.sarif"})
            assert not result.isError
            assert json.loads(result.content[0].text)["counts_by_level"] == {"error": 1}
            result = await session.call_tool("read_document", {"path": "report.pdf", "max_chars": 1000})
            assert not result.isError
            assert "Reader evidence" in json.loads(result.content[0].text)["text"]
            for path in ("https://example.invalid/report.pdf", "../outside.pdf"):
                result = await session.call_tool("read_document", {"path": path})
                assert result.isError, "MCP must reject inputs outside the workspace"
    assert {path.name for path in root.iterdir()} == {
        "report.pdf", "report.docx", "report.xlsx", "report.pptx", "long.docx", "findings.sarif",
    }, "reader must not leave converted files in the workspace"


if __name__ == "__main__":
    with tempfile.TemporaryDirectory(prefix="report-reader-smoke-") as temp:
        asyncio.run(asyncio.wait_for(smoke(Path(temp)), timeout=150))
    print("Report reader: four real document conversions and MCP stdio calls passed.")
